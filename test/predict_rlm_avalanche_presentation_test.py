from pathlib import Path

import pytest
from openpyxl import load_workbook
from predict_rlm import File

from examples.predict_rlm_avalanche_presentation.config import FEEDBACK_WORKBOOK_PATH
from examples.predict_rlm_avalanche_presentation.flow import compose_product_review
from examples.predict_rlm_avalanche_presentation.schema import (
    AccountExposure,
    EvidenceReference,
    Risk,
    RiskReport,
    Theme,
    ThemeReport,
)
from examples.predict_rlm_avalanche_presentation.util import publish_review_pack_files


def test_feedback_workbook_has_valid_joinable_source_data():
    workbook = load_workbook(FEEDBACK_WORKBOOK_PATH, read_only=True, data_only=True)

    assert workbook.sheetnames == ["Feedback", "Accounts", "Roadmap"]
    feedback_rows = list(workbook["Feedback"].iter_rows(values_only=True))
    account_rows = list(workbook["Accounts"].iter_rows(values_only=True))
    roadmap_rows = list(workbook["Roadmap"].iter_rows(values_only=True))
    workbook.close()

    assert feedback_rows[0] == (
        "feedback_id",
        "account_id",
        "text",
        "product_area",
        "channel",
        "sentiment",
    )
    assert len(feedback_rows) == 101
    assert len(account_rows) == 25
    assert len(roadmap_rows) == 41

    account_ids = {row[0] for row in account_rows[1:]}
    roadmap_areas = {row[0] for row in roadmap_rows[1:]}
    assert {row[1] for row in feedback_rows[1:]} <= account_ids
    assert {row[3] for row in feedback_rows[1:]} <= roadmap_areas


def test_compose_product_review_validates_deterministic_totals():
    evidence = EvidenceReference(
        feedback_id="fb-001",
        account_id="ent-001",
        sheet="Feedback",
        row_number=2,
        excerpt="Export jobs time out.",
    )
    themes = ThemeReport(
        feedback_rows_analyzed=100,
        themes=[
            Theme(
                name="Reliable exports",
                summary="Large exports must complete reliably.",
                feedback_count=2,
                account_count=2,
                segment_counts={"enterprise": 1, "mid_market": 1},
                evidence=[
                    evidence,
                    EvidenceReference(
                        feedback_id="fb-002",
                        account_id="mid-001",
                        sheet="Feedback",
                        row_number=3,
                        excerpt="Failed exports must be retried.",
                    ),
                ],
            )
        ],
    )
    exposure = AccountExposure(
        account_id="ent-001",
        account_name="Northstar Bank",
        customer_segment="enterprise",
        arr_usd=240_000,
        renewal_date="2026-08-15",
    )
    risk = Risk(
        issue="Export failures block reporting",
        severity="high",
        rationale="Month-end reporting cannot complete.",
        product_areas=["data_exports"],
        affected_accounts=[exposure],
        arr_at_risk_usd=240_000,
        evidence=[evidence],
    )
    risks = RiskReport(feedback_rows_analyzed=100, risks=[risk])

    review = compose_product_review.fn(themes, risks)

    assert review.feedback_rows_analyzed == 100
    assert review.themes == themes.themes
    assert review.risks == risks.risks

    invalid_risks = RiskReport(
        feedback_rows_analyzed=100,
        risks=[risk.model_copy(update={"arr_at_risk_usd": 1})],
    )
    with pytest.raises(ValueError, match="ARR total"):
        compose_product_review.fn(themes, invalid_risks)


def test_publish_review_pack_copies_named_artifacts(tmp_path: Path):
    workbook_source = tmp_path / "generated.xlsx"
    brief_source = tmp_path / "generated.docx"
    workbook_source.write_bytes(b"workbook")
    brief_source.write_bytes(b"brief")

    result = publish_review_pack_files(
        File(path=str(workbook_source)),
        File(path=str(brief_source)),
        destination=tmp_path / "published",
    )

    assert Path(result.workbook_path).read_bytes() == b"workbook"
    assert Path(result.brief_path).read_bytes() == b"brief"
    assert Path(result.workbook_path).name == "product_review.xlsx"
    assert Path(result.brief_path).name == "executive_brief.docx"
